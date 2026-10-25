"""Export the exercise library to an Excel workbook for a manual content-review pass. Rerun it any time the data
changes; it overwrites the file.

Run from the project root (Windows):
    backend\\.venv\\Scripts\\python.exe backend\\scripts\\export_exercise_review.py
or from backend/:
    .venv\\Scripts\\python.exe scripts\\export_exercise_review.py

Writes content/exercise_review_export.xlsx (gitignored: it's rebuilt from the database, so there's no point
versioning a binary copy). It reads the working database (backend/app/db/homebase.db, or HOMEBASE_DB_PATH) read-only.

Sheet "Exercises": one row per exercise, sorted by muscle group then name.
Sheet "Quiz Logic": every quiz question (wording and options read live from the database), the consent gate and the
calendar choice, with the column each is stored in, how the app uses it, and its safety / privacy notes. The script
stops with an error if a question or a quiz_profiles column isn't covered, so the sheet can't silently go stale.
"""

import json
import sys
from pathlib import Path

if __package__ in (None, ""):   # run as a script: make `import app...` work
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import config
from app.db import connect_readonly

OUT = config.CONTENT_DIR / "exercise_review_export.xlsx"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")   # blue and white (AGENTS.md)
HEADER_FONT = Font(name="Times New Roman", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Times New Roman")

EXERCISE_COLUMNS = [   # (header, width)
    ("name", 34), ("muscle_group", 13), ("equipment_needed", 20), ("difficulty", 12),
    ("injury draft areas", 26), ("injury reasoning (AI-drafted, pending PT review)", 90), ("source_dataset", 20),
    ("exercise_id", 30),
]


def _words(values: list[str]) -> str:
    return ", ".join(values)


def exercise_rows() -> list[list[str]]:
    conn = connect_readonly()
    try:
        rows = conn.execute(
            "SELECT id, name, muscle_group, equipment_needed, difficulty, injury_contraindications, injury_rationale, "
            "source_dataset, contraindication_status FROM exercises ORDER BY muscle_group, name COLLATE NOCASE").fetchall()
    finally:
        conn.close()
    out = []
    for r in rows:
        tags = json.loads(r["injury_contraindications"])
        rationale = json.loads(r["injury_rationale"] or "{}")
        reasoning = "\n".join(f"{t}: {rationale[t]['reasoning']}" for t in tags if t in rationale)
        if r["contraindication_status"] == "untagged":          # nobody has looked at it yet
            areas, reasoning = "(not checked yet)", "No injury cautions drafted yet."
        elif not tags:                                           # looked at, and judged to need none
            areas, reasoning = "none", "Drafted with no injury cautions (none judged plausible)."
        else:
            areas = _words(tags)
        out.append([r["name"], r["muscle_group"], _words(json.loads(r["equipment_needed"])), r["difficulty"],
                    areas, reasoning, r["source_dataset"], r["id"]])
    return out


# ---------------------------------------------------------------- sheet 2: Quiz Logic
#
# Traced from the code on 2026-10-25 (after the quiz and plan rewrite): content/quiz/quiz_questions.yaml, app/users.py (QUIZ_COLUMNS, save_quiz,
# reset_quiz), app/db/schema.sql (quiz_profiles), app/services/program_generator.py, app/agents/context.py and the
# agents, app/mcp_server/progress.py, app/services/workout_review.py, cardio_estimate.py, progress_marker.py,
# app/calendar.py, and the quiz page (frontend/src/pages/QuizPage.tsx). Wording and options come live from the
# database; check_quiz_coverage() refuses to write the sheet if a question or a quiz_profiles column has no row here.
AGENT_PROFILE = ("part of the PII-free agent profile (app/agents/context.py: just the goal and injury flags), which "
                 "the Workout Builder's model can read through its get_user_profile tool")
QUIZ_LOGIC: dict[str, dict] = {
    "goal": {
        "db": "quiz_profiles.goal",
        "used": "Picks the week (always 4 training days): Getting stronger = 3 strength days (each Upper OR Lower) + "
                "1 cardio, repeating every 2 weeks; Losing weight = 2 strength + 2 cardio in a 4-week pattern; Feeling "
                "good = 2 mobility/core + 1 strength + 1 cardio in an 8-week pattern. With experience level it picks "
                "the goal_protocols row (sets, reps, rest, session length). It sets how weights go up: Getting "
                "stronger +2.5 lb (arms) / +5 lb (legs) each time a day repeats; Losing weight +5 lb every 4 weeks; "
                "Feeling good none. It decides whether weight is asked, the goal check's goal-specific check, and "
                "whether pound figures appear (Losing weight only).",
        "safety": "Not identifying. Sent to models: in the Program Architect's prompt, in the Coach's "
                  "check_goal_achievement result, and " + AGENT_PROFILE + ".",
    },
    "experience_level": {
        "db": "quiz_profiles.experience_level",
        "used": "Picks the goal_protocols row (goal x level), caps exercise difficulty (beginner: beginner moves only; "
                "intermediate: up to intermediate; advanced: any), keeps beginners' plans low-impact (no jumping), and "
                "sets the default starting weights: beginner none (bodyweight only), intermediate 10 lb arms / 20 lb "
                "legs, advanced 20 / 40. Always editable.",
        "safety": "Sent to a model in the Program Architect's prompt. Not in the agent profile.",
    },
    "equipment": {
        "db": "quiz_profiles.equipment (JSON list; bodyweight is always included)",
        "used": "Filters exercises in the plan: bodyweight (and a sturdy chair) is always allowed; dumbbells, "
                "kettlebells, resistance bands and a mat add what needs them; 'A full gym' removes the equipment "
                "filter (and is the only way bar moves like pull-ups appear). Exercises that use weights get a weight "
                "field. Build Your Own uses the equipment picked on that page, not this answer.",
        "safety": "Not in the agent profile; never sent to a model (filters are applied in code).",
    },
    "injuries": {
        "db": "quiz_profiles.injury_flags (JSON list; 'None' stores [])",
        "used": "Removes every exercise tagged unsafe for any flagged area (the SQL filter in search_exercises, used "
                "by the plan generator, the Workout Builder and the MCP tool) and forces checked_only (only exercises "
                "with drafted injury data). The plan lists what was left out and why, and stays low-impact. Applies on "
                "top of whichever goal is picked. Weights still go up by the goal's rule (decided 2026-10-25).",
        "safety": "checked_only is set by code and the model can't loosen it. The flags (never the reasons) are "
                  + AGENT_PROFILE + "; the Program Architect sees only a yes/no 'low-impact injury-safe'.",
    },
    "injuries:free_text": {
        "label": "Anything else we should know? (free text under the injury question)",
        "format": "Free text, optional, up to 1,000 characters",
        "db": "quiz_profiles.injury_notes",
        "used": "Stored only. No code reads it after saving.",
        "safety": "Never sent to any model (on the excluded list in app/agents/context.py, with username, first name "
                  "and email). Deleted with the account.",
    },
    "pregnancy_status": {
        "db": "quiz_profiles.pregnancy_status, plus quiz_profiles.needs_clinical_review (computed by the database: "
              "1 for pregnant or postpartum)",
        "used": "Pregnant or postpartum: no plan is generated, the quiz sends them to /pregnancy-support, the "
                "calendar stays empty, Build Your Own and the calendar's day choices are blocked, and every AI route "
                "refuses (403, or a code-only reply for Ask the Basics after the red-flag check). They can still log "
                "cardio in their own words (no library picker), with no calorie estimates and no progress marker.",
        "safety": "Triggers needs_clinical_review, which blocks every AI feature, so no model is ever called for "
                  "them. Not in the agent profile; never sent to any model.",
    },
    "weight": {
        "db": "quiz_profiles.weight (lb; only stored when the goal is losing weight)",
        "used": "Asked only when the goal is Losing weight. Calorie math done in code: the Workout Reviewer's estimate, "
                "cardio log estimates, calories on done calendar days and the progress marker (kcal = MET x kg x "
                "hours). Other goals have no weight, so they see no calorie numbers.",
        "safety": "Never sent to any model (not in the agent profile since 2026-10-25; on its excluded list). The "
                  "database refuses a weight on any other goal. No calorie numbers while pregnant or postpartum.",
    },
    "goal_weight": {
        "db": "quiz_profiles.goal_weight (lb, optional; only stored when the goal is losing weight)",
        "used": "Asked only when the goal is Losing weight. Used only on the quiz page: a gentle CDC-cited note on "
                "roughly how many weeks the change takes at 1 to 2 lb a week (never a block). No backend logic reads "
                "it.",
        "safety": "Not in the agent profile; never sent to any model.",
    },
}
# quiz_profiles columns that aren't quiz answers: the consent gate and what the app sets around the quiz.
QUIZ_EXTRA_ROWS = [
    {"label": "Medical disclaimer (checkbox before the health questions)",
     "format": "Required checkbox: \"I have read and understand this disclaimer...\"; the version shown is recorded",
     "db": "quiz_profiles.consent_accepted, consent_accepted_at (UTC), consent_version",
     "used": "The consent gate: the health questions stay hidden until it's ticked, the API refuses the quiz without "
             "it (422), and the database refuses a profile without it (CHECK). \"Reset my quiz\" deletes it, so a "
             "retake asks again.",
     "safety": "Records that the disclaimer was accepted and when. Never sent to a model."},
    {"label": "How do you want to populate your calendar? (asked right after the quiz; changeable in Settings)",
     "format": "Single-select: Use my quiz plan (generated, the default) | Let me review it first (review) | "
               "I'll build it myself (self)",
     "db": "quiz_profiles.calendar_mode",
     "used": "generated: the plan is laid out on the calendar. review: nothing is laid out until the user reviews "
             "the plan and optionally swaps exercises. self: days are open and filled with Build Your Own or cardio. "
             "Kept when the quiz is retaken.",
     "safety": "Blocked for pregnant / postpartum users (403). Not sent to a model."},
    {"label": "(set by the app) review-mode bookkeeping", "format": "Not asked",
     "db": "quiz_profiles.plan_reviewed, quiz_profiles.plan_swaps (JSON)",
     "used": "Whether a review-mode plan has been finalized, and the user's exercise swaps. Swaps are only accepted "
             "from the same injury-, equipment- and level-filtered pool. Both reset on every quiz save.",
     "safety": "Swaps can't get around the injury filter (checked when saved and again when the plan is built)."},
    {"label": "(set by the app) plan start", "format": "Not asked", "db": "quiz_profiles.plan_started_on",
     "used": "Week 1 of the plan: which week of the 2-, 4- or 8-week pattern a date falls in, and how far weights have "
             "gone up. Week 1 is the calibration week (log what you actually did; later suggestions build on it). "
             "Reset to today on every quiz save.",
     "safety": "Weight suggestions are always editable, and going up applies to everyone by the goal's rule, "
               "including people with an injury flag (decided 2026-10-25)."},
    {"label": "(set by the app) record keeping", "format": "Not asked",
     "db": "quiz_profiles.id, user_id, created_at, updated_at",
     "used": "created_at (the first quiz save) starts the goal check's clock: under 7 days it says it's too early to "
             "judge. user_id links the profile to the account.",
     "safety": "user_id is never sent to a model. The whole row is deleted with the account, or by \"Reset my quiz\"."},
]
QUIZ_COLUMNS = [("Quiz question", 42), ("Answer options / format", 48), ("Database field", 34),
                ("How it's used downstream", 70), ("Safety / privacy notes", 60)]


def _format(q, settings: dict) -> str:
    options = json.loads(q["options"])
    req = "required" if q["required"] else "optional"
    if q["question_type"] in ("multi_select", "single_select"):
        kind = "Multi-select" if q["question_type"] == "multi_select" else "Single-select"
        excl = f"; '{settings['exclusive']}' clears the others" if settings.get("exclusive") else ""
        return f"{kind} ({req}{excl}): " + " | ".join(f"{o['label']} ({o['value']})" for o in options)
    if q["question_type"] == "number":
        if settings.get("input") == "feet_inches":
            return f"Feet + inches ({req}), stored as total inches, {settings.get('min')} to {settings.get('max')} in"
        return f"Number ({req}), {settings.get('min')} to {settings.get('max')} {settings.get('unit', '')}".strip()
    if q["question_type"] == "date":
        return f"Date ({req})"
    return f"Free text ({req}), up to 500 characters"


def check_quiz_coverage(conn, questions) -> None:
    """Fail loudly if the quiz or the table changed and this sheet would be missing something."""
    ids = {q["id"] for q in questions}
    missing_q = ids - {k for k in QUIZ_LOGIC if ":" not in k}
    stale_q = {k for k in QUIZ_LOGIC if ":" not in k} - ids
    import re
    described = set(re.findall(r"[a-z_]+", " ".join(r["db"] for r in [*QUIZ_LOGIC.values(), *QUIZ_EXTRA_ROWS])))
    columns = [r[1] for r in conn.execute("PRAGMA table_xinfo(quiz_profiles)")]
    missing_c = [c for c in columns if c not in described]   # whole names only
    if missing_q or stale_q or missing_c:
        raise SystemExit(f"Quiz Logic sheet is out of date. Questions without a row: {sorted(missing_q)}; rows for "
                         f"questions that no longer exist: {sorted(stale_q)}; quiz_profiles columns not covered: "
                         f"{missing_c}. Update QUIZ_LOGIC in {Path(__file__).name}.")


def quiz_rows() -> list[list[str]]:
    conn = connect_readonly()
    try:
        questions = conn.execute("SELECT * FROM quiz_questions ORDER BY order_index").fetchall()
        check_quiz_coverage(conn, questions)
    finally:
        conn.close()
    rows = []
    for q in questions:
        settings = json.loads(q["settings"])
        logic = QUIZ_LOGIC[q["id"]]
        stage = " [health question: shown after the disclaimer]" if settings.get("health") else ""
        rows.append([q["prompt"] + stage, _format(q, settings), logic["db"], logic["used"], logic["safety"]])
        if settings.get("allow_free_text"):
            extra = QUIZ_LOGIC[f"{q['id']}:free_text"]
            rows.append([f"{settings.get('free_text_prompt') or extra['label']} (free text under \"{q['prompt']}\")",
                         extra["format"], extra["db"], extra["used"], extra["safety"]])
    rows += [[r["label"], r["format"], r["db"], r["used"], r["safety"]] for r in QUIZ_EXTRA_ROWS]
    return rows


def write_sheet(ws, columns: list[tuple[str, int]], rows: list[list], wrap: set[int] = frozenset()) -> None:
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column - 1 in wrap)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def build() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Exercises"
    write_sheet(ws, EXERCISE_COLUMNS, exercise_rows(), wrap={5})
    write_sheet(wb.create_sheet("Quiz Logic"), QUIZ_COLUMNS, quiz_rows(), wrap={0, 1, 2, 3, 4})
    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(OUT)
    except PermissionError:
        raise SystemExit(f"Couldn't write {OUT.name}: it's probably open in Excel. Close it and run this again.")
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
